import base64
import hashlib
import io
import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import date, datetime
from typing import Any

import fitz
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openai import OpenAI


# ============================================================
# App configuration
# ============================================================
st.set_page_config(
    page_title="EBOC Guideline Update Assistant v2",
    page_icon="",
    layout="wide",
)

APP_TITLE = "EBOC Guideline Update Assistant v2"
APP_SUBTITLE = "Evidence Validation Edition"
MODEL_NAME = "gpt-5.5"
NCBI_TOOL_NAME = "eboc_guideline_update_assistant"
BANNER_PATH = "DCMC Banner.jpg"
LOGO_PATH = "EBOC Logo.JPG"

PROJECT_TYPES = (
    "Revision or Update of Existing Guideline",
    "New Guideline Development",
)


# ============================================================
# Styling and branding
# ============================================================
st.markdown(
    """
    <style>
    .block-container {
        padding-top: 1.2rem;
        padding-bottom: 3rem;
    }
    .eboc-banner {
        width: 100%;
        max-height: 220px;
        object-fit: cover;
        object-position: center;
        border-radius: 10px;
        margin-bottom: 0.9rem;
    }
    .eboc-title {
        font-size: 2.15rem;
        font-weight: 750;
        line-height: 1.12;
        margin: 0.25rem 0 0 0;
    }
    .eboc-subtitle {
        font-size: 1.12rem;
        font-weight: 600;
        color: #4b5563;
        margin-top: 0.25rem;
    }
    .workflow-box {
        padding: 0.65rem 0.9rem;
        border: 1px solid #d8dee8;
        border-radius: 9px;
        background: #f8fafc;
        margin: 0.55rem 0 1.15rem 0;
        font-size: 0.94rem;
    }
    .status-chip {
        display: inline-block;
        padding: 0.38rem 0.62rem;
        border-radius: 999px;
        font-size: 0.88rem;
        font-weight: 600;
        margin: 0.12rem 0;
        line-height: 1.2;
    }
    .status-ready {
        background: #e8f5e9;
        color: #1b5e20;
        border: 1px solid #a5d6a7;
    }
    .status-missing {
        background: #fff4e5;
        color: #8a4b08;
        border: 1px solid #f3c78a;
    }
    .status-neutral {
        background: #f1f5f9;
        color: #475569;
        border: 1px solid #cbd5e1;
    }
    .confidence-badge {
        display: inline-flex;
        align-items: center;
        gap: 0.55rem;
        padding: 0.55rem 0.8rem;
        border-radius: 10px;
        font-weight: 700;
        margin: 0.25rem 0 0.8rem 0;
        border: 1px solid;
    }
    .confidence-dot {
        width: 16px;
        height: 16px;
        border-radius: 50%;
        display: inline-block;
    }
    .validation-card {
        border: 1px solid #d8dee8;
        border-radius: 9px;
        padding: 0.8rem 0.95rem;
        margin-bottom: 0.65rem;
        background: #ffffff;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


def image_as_data_uri(path: str) -> str | None:
    """Return a local image as a data URI for controlled display sizing."""
    if not os.path.exists(path):
        return None
    with open(path, "rb") as image_file:
        encoded = base64.b64encode(image_file.read()).decode("utf-8")
    extension = os.path.splitext(path)[1].lower()
    mime = "image/png" if extension == ".png" else "image/jpeg"
    return f"data:{mime};base64,{encoded}"


def render_branding() -> None:
    banner_uri = image_as_data_uri(BANNER_PATH)
    if banner_uri:
        st.markdown(
            f'<img src="{banner_uri}" class="eboc-banner" alt="Dell Children\'s banner">',
            unsafe_allow_html=True,
        )

    logo_col, title_col = st.columns([1, 5], vertical_alignment="center")
    with logo_col:
        if os.path.exists(LOGO_PATH):
            st.image(LOGO_PATH, width=165)
    with title_col:
        st.markdown(f'<div class="eboc-title">{APP_TITLE}</div>', unsafe_allow_html=True)
        st.markdown(f'<div class="eboc-subtitle">{APP_SUBTITLE}</div>', unsafe_allow_html=True)

    st.markdown(
        """
        <div class="workflow-box">
        <b>Workflow:</b> 1. Guideline details &nbsp;→&nbsp; 2. Upload documents
        &nbsp;→&nbsp; 3. Evidence validation &nbsp;→&nbsp; 4. Readiness check
        &nbsp;→&nbsp; 5. Generate &nbsp;→&nbsp; 6. Preview and download
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# Secrets and API configuration
# ============================================================
def get_secret(name: str) -> str | None:
    try:
        value = st.secrets.get(name, None)
        return str(value).strip() if value else None
    except Exception:
        return None


def get_api_key() -> str | None:
    """Read the OpenAI key from Streamlit Secrets, with a test fallback."""
    api_key = get_secret("OPENAI_API_KEY")
    if not api_key:
        api_key = st.sidebar.text_input(
            "OpenAI API key",
            type="password",
            help=(
                "Temporary option for testing. In the deployed app, store the key "
                "in Streamlit Secrets rather than in GitHub."
            ),
        )
    return api_key


# ============================================================
# Document reading
# ============================================================
def read_docx_bytes(file_bytes: bytes) -> str:
    """Read paragraphs and tables from a Word document."""
    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts)


def read_pdf_bytes(file_bytes: bytes) -> str:
    """Read extractable text from a PDF."""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = [page.get_text() for page in doc]
    return "\n".join(pages)


def read_uploaded_file(uploaded_file: Any) -> str:
    """Read a Streamlit UploadedFile without consuming it permanently."""
    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".docx"):
        return read_docx_bytes(file_bytes)
    if filename.endswith(".pdf"):
        return read_pdf_bytes(file_bytes)

    raise ValueError("Unsupported file type. Upload a DOCX or PDF file.")


# ============================================================
# General text helpers
# ============================================================
def normalize_title(title: str | None) -> str:
    if not title:
        return ""
    normalized = re.sub(r"<[^>]+>", " ", title)
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized.lower())
    return re.sub(r"\s+", " ", normalized).strip()


def safe_json_loads(text: str) -> dict[str, Any]:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        parsed = json.loads(cleaned)
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
        if not match:
            return {}
        try:
            parsed = json.loads(match.group(0))
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}




def validation_context_signature(
    guideline_name: str,
    focused_topic: str,
    evidence_question: str,
    literature_file: Any,
) -> str:
    """Fingerprint the inputs that determine an evidence validation search."""
    literature_digest = ""
    if literature_file is not None:
        literature_digest = hashlib.sha256(literature_file.getvalue()).hexdigest()

    payload = "\n".join(
        [
            guideline_name.strip(),
            focused_topic.strip(),
            evidence_question.strip(),
            literature_digest,
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()

def current_year() -> int:
    return date.today().year


def format_display_date() -> str:
    return f"{date.today().month}/{date.today().day}/{date.today().year}"


# ============================================================
# Word formatting
# ============================================================
def add_hyperlink(paragraph, url: str, text: str | None = None) -> None:
    if text is None:
        text = url

    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)

    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_markdown_runs(paragraph, text: str) -> None:
    """Convert basic bold markdown and URLs into Word runs."""
    url_pattern = r"(https?://[^\s]+)"
    parts = re.split(url_pattern, text)

    for part in parts:
        if not part:
            continue

        if re.match(url_pattern, part):
            clean_url = part.rstrip(".,);]")
            trailing = part[len(clean_url):]
            add_hyperlink(paragraph, clean_url, clean_url)
            if trailing:
                paragraph.add_run(trailing)
            continue

        bold_parts = re.split(r"(\*\*.*?\*\*)", part)
        for bold_part in bold_parts:
            if not bold_part:
                continue
            if bold_part.startswith("**") and bold_part.endswith("**"):
                run = paragraph.add_run(bold_part[2:-2])
                run.bold = True
            else:
                paragraph.add_run(bold_part)


def is_bullet_line(line: str) -> bool:
    stripped = line.strip()
    return any(stripped.startswith(prefix) for prefix in ("- ", "● ", "* ", "• "))


def is_numbered_line(line: str) -> bool:
    return re.match(r"^\d+\.\s+", line.strip()) is not None


def clean_bullet_text(line: str) -> str:
    stripped = line.strip()
    for prefix in ("- ", "● ", "* ", "• "):
        if stripped.startswith(prefix):
            return stripped[len(prefix):].strip()
    return stripped


def clean_numbered_text(line: str) -> str:
    return re.sub(r"^\d+\.\s+", "", line.strip())


def set_normal_spacing(paragraph) -> None:
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08


def confidence_color(level: str) -> tuple[str, RGBColor, str]:
    if level == "High":
        return "Green", RGBColor(46, 125, 50), "2E7D32"
    if level == "Moderate":
        return "Yellow", RGBColor(245, 158, 11), "F59E0B"
    if level == "Low":
        return "Red", RGBColor(198, 40, 40), "C62828"
    return "Not identified", RGBColor(71, 85, 105), "475569"


def add_word_confidence_badge(doc: Document, level: str) -> None:
    label, color, _ = confidence_color(level)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)

    dot = paragraph.add_run("● ")
    dot.font.color.rgb = color
    dot.bold = True
    dot.font.size = Pt(14)

    rating = paragraph.add_run(f"{level} confidence")
    rating.bold = True
    rating.font.size = Pt(11)

    descriptor = paragraph.add_run(f"  |  Stoplight indicator: {label}")
    descriptor.font.size = Pt(10)


def create_styled_docx(
    card_text: str,
    guideline_name: str,
    topic: str,
    confidence_level: str,
) -> io.BytesIO:
    """Create a clean physician facing Word document."""
    doc = Document()
    styles = doc.styles

    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True
    styles["Heading 3"].font.name = "Calibri"
    styles["Heading 3"].font.size = Pt(11)
    styles["Heading 3"].font.bold = True

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run(
        f"{format_display_date()}\n{guideline_name}\n{topic}"
    )
    header_run.bold = True
    header_run.font.size = Pt(10)
    doc.add_paragraph("")

    confidence_badge_added = False

    for raw_line in card_text.splitlines():
        line = raw_line.strip()

        if not line:
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line.replace("# ", "", 1).strip(), level=1)
            continue

        if line.startswith("## "):
            heading_text = line.replace("## ", "", 1).strip()
            doc.add_heading(heading_text, level=2)
            if "GRADE" in heading_text and "Confidence" in heading_text:
                add_word_confidence_badge(doc, confidence_level)
                confidence_badge_added = True
            continue

        if line.startswith("### "):
            doc.add_heading(line.replace("### ", "", 1).strip(), level=3)
            continue

        if is_bullet_line(line):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(paragraph, clean_bullet_text(line))
            set_normal_spacing(paragraph)
            continue

        if is_numbered_line(line):
            paragraph = doc.add_paragraph(style="List Number")
            add_markdown_runs(paragraph, clean_numbered_text(line))
            set_normal_spacing(paragraph)
            continue

        if line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            paragraph.paragraph_format.space_after = Pt(6)
            add_markdown_runs(paragraph, line.replace("> ", "", 1).strip())
            continue

        if confidence_badge_added and re.match(r"^\*\*Rating:\*\*", line, flags=re.IGNORECASE):
            continue

        paragraph = doc.add_paragraph()
        add_markdown_runs(paragraph, line)
        set_normal_spacing(paragraph)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


# ============================================================
# OpenAI helpers
# ============================================================
def create_openai_client(api_key: str) -> OpenAI:
    return OpenAI(api_key=api_key)


def extract_literature_inventory(
    client: OpenAI,
    literature_text: str,
) -> dict[str, Any]:
    """Extract citation fingerprints from uploaded literature evidence."""
    instructions = """
You extract citation metadata from supplied literature evidence.
Use only the supplied text. Do not invent missing information.
Return valid JSON only, with this exact top level structure:
{
  "records": [
    {
      "title": null,
      "year": null,
      "pmid": null,
      "doi": null,
      "evidence_type": null
    }
  ]
}
Use null when a value is unavailable. Include every identifiable cited source.
"""

    response = client.responses.create(
        model=MODEL_NAME,
        instructions=instructions,
        input=literature_text,
        store=False,
    )
    parsed = safe_json_loads(response.output_text)
    records = parsed.get("records", [])
    if not isinstance(records, list):
        records = []

    clean_records: list[dict[str, Any]] = []
    for record in records:
        if not isinstance(record, dict):
            continue
        year_value = record.get("year")
        try:
            year_value = int(year_value) if year_value else None
        except (TypeError, ValueError):
            year_value = None
        clean_records.append(
            {
                "title": record.get("title"),
                "year": year_value,
                "pmid": str(record.get("pmid") or "").strip() or None,
                "doi": str(record.get("doi") or "").strip().lower() or None,
                "evidence_type": record.get("evidence_type"),
            }
        )

    if not clean_records:
        clean_records = fallback_literature_inventory(literature_text)

    years = [
        record["year"]
        for record in clean_records
        if isinstance(record.get("year"), int)
        and 1900 <= record["year"] <= current_year() + 1
    ]

    return {
        "records": clean_records,
        "newest_year": max(years) if years else None,
    }


def fallback_literature_inventory(text: str) -> list[dict[str, Any]]:
    """Create a minimal inventory if model based extraction is unavailable."""
    pmids = sorted(set(re.findall(r"\bPMID\s*:?\s*(\d{6,9})\b", text, flags=re.IGNORECASE)))
    dois = sorted(
        set(
            match.rstrip(".,;)").lower()
            for match in re.findall(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", text, flags=re.IGNORECASE)
        )
    )
    years = [
        int(year)
        for year in re.findall(r"\b(?:19|20)\d{2}\b", text)
        if 1900 <= int(year) <= current_year() + 1
    ]

    records: list[dict[str, Any]] = []
    for pmid in pmids:
        records.append({"title": None, "year": None, "pmid": pmid, "doi": None, "evidence_type": None})
    for doi in dois:
        records.append({"title": None, "year": None, "pmid": None, "doi": doi, "evidence_type": None})

    if not records and years:
        records.append({"title": None, "year": max(years), "pmid": None, "doi": None, "evidence_type": None})

    return records


def build_pubmed_query(
    client: OpenAI,
    guideline_name: str,
    focused_topic: str,
    evidence_question: str,
) -> str:
    """Create a focused PubMed query without date or journal limits."""
    instructions = """
Create one valid PubMed search query for a pediatric clinical evidence validation search.
Return only the query text with no explanation and no code block.
Include the core condition, population, intervention or exposure, and outcome concepts when supported by the question.
Use PubMed field tags and Boolean operators where helpful.
Include pediatric terms unless the question already clearly limits the population.
Do not add publication date limits.
Do not add journal filters.
Keep the query under 900 characters.
"""

    input_text = f"""
Guideline name: {guideline_name}
Focused topic: {focused_topic}
Evidence question: {evidence_question}
"""

    try:
        response = client.responses.create(
            model=MODEL_NAME,
            instructions=instructions,
            input=input_text,
            store=False,
        )
        query = response.output_text.strip()
        query = re.sub(r"^```(?:text)?\s*", "", query, flags=re.IGNORECASE)
        query = re.sub(r"\s*```$", "", query)
        if query:
            return query[:900]
    except Exception:
        pass

    condition = re.sub(r"[^A-Za-z0-9 ]+", " ", guideline_name).strip()
    topic = re.sub(r"[^A-Za-z0-9 ]+", " ", focused_topic).strip()
    return (
        f'("{condition}"[Title/Abstract] OR "{topic}"[Title/Abstract]) '
        "AND (pediatric*[Title/Abstract] OR child*[Title/Abstract] "
        "OR infant*[Title/Abstract] OR adolescent*[Title/Abstract])"
    )


# ============================================================
# NCBI PubMed helpers
# ============================================================
def ncbi_request(
    endpoint: str,
    parameters: dict[str, Any],
    email: str,
    timeout: int = 40,
) -> bytes:
    base_url = f"https://eutils.ncbi.nlm.nih.gov/entrez/eutils/{endpoint}"
    common = {
        "tool": NCBI_TOOL_NAME,
        "email": email,
    }
    query_parameters = {**common, **parameters}
    url = f"{base_url}?{urllib.parse.urlencode(query_parameters)}"

    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": f"{NCBI_TOOL_NAME}/2.0 ({email})",
            "Accept": "application/json, application/xml, text/xml, */*",
        },
    )

    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.read()
    except urllib.error.HTTPError as error:
        raise RuntimeError(f"NCBI returned HTTP {error.code}.") from error
    except urllib.error.URLError as error:
        raise RuntimeError(f"NCBI connection failed: {error.reason}") from error


def pubmed_search(query: str, email: str, retmax: int = 15) -> list[str]:
    raw = ncbi_request(
        "esearch.fcgi",
        {
            "db": "pubmed",
            "term": query,
            "retmode": "json",
            "retmax": retmax,
            "sort": "pub date",
        },
        email=email,
    )
    payload = json.loads(raw.decode("utf-8"))
    return payload.get("esearchresult", {}).get("idlist", [])


def element_text(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(element.itertext()).strip()


def extract_publication_year(article: ET.Element) -> int | None:
    year_paths = (
        ".//JournalIssue/PubDate/Year",
        ".//ArticleDate/Year",
        ".//DateCompleted/Year",
        ".//DateRevised/Year",
    )
    for path in year_paths:
        text = article.findtext(path)
        if text and text.isdigit():
            year = int(text)
            if 1900 <= year <= current_year() + 1:
                return year

    medline_date = article.findtext(".//JournalIssue/PubDate/MedlineDate") or ""
    match = re.search(r"\b(?:19|20)\d{2}\b", medline_date)
    return int(match.group(0)) if match else None


def classify_evidence_type(publication_types: list[str], journal: str) -> str:
    joined = " | ".join(publication_types).lower()
    journal_lower = journal.lower()

    if "guideline" in joined or "practice guideline" in joined:
        return "Clinical guideline"
    if "meta-analysis" in joined or "systematic review" in joined or "cochrane" in journal_lower:
        return "Systematic review or meta analysis"
    if "randomized controlled trial" in joined or "controlled clinical trial" in joined:
        return "Randomized trial"
    if "consensus development conference" in joined:
        return "Consensus statement"
    if any(term in joined for term in ("observational study", "cohort", "case-control")):
        return "Observational study"
    if "review" in joined:
        return "Review"
    return "Other evidence"


def classify_source_group(journal: str, iso_abbreviation: str) -> str:
    combined = f"{journal} {iso_abbreviation}".lower()
    if "jama pediatr" in combined or "jama pediatrics" in combined:
        return "JAMA Pediatrics"
    if any(
        name in combined
        for name in ("pediatrics", "hospital pediatrics", "pediatrics in review")
    ):
        return "AAP publication"
    if "cochrane database" in combined:
        return "Cochrane review"
    return "Other PubMed indexed source"


def pubmed_fetch(pmids: list[str], email: str) -> list[dict[str, Any]]:
    if not pmids:
        return []

    raw = ncbi_request(
        "efetch.fcgi",
        {
            "db": "pubmed",
            "id": ",".join(pmids),
            "retmode": "xml",
        },
        email=email,
    )

    root = ET.fromstring(raw)
    records: list[dict[str, Any]] = []

    for pubmed_article in root.findall(".//PubmedArticle"):
        citation = pubmed_article.find("MedlineCitation")
        article = pubmed_article.find(".//Article")
        if citation is None or article is None:
            continue

        pmid = element_text(citation.find("PMID"))
        title = element_text(article.find("ArticleTitle"))
        journal = element_text(article.find("Journal/Title"))
        iso_abbreviation = element_text(article.find("Journal/ISOAbbreviation"))
        year = extract_publication_year(pubmed_article)

        authors: list[str] = []
        for author in article.findall("AuthorList/Author"):
            collective = element_text(author.find("CollectiveName"))
            if collective:
                authors.append(collective)
                continue
            last_name = element_text(author.find("LastName"))
            initials = element_text(author.find("Initials"))
            author_name = " ".join(part for part in (last_name, initials) if part)
            if author_name:
                authors.append(author_name)

        abstract_parts: list[str] = []
        for abstract_text in article.findall("Abstract/AbstractText"):
            section_label = abstract_text.attrib.get("Label")
            text = element_text(abstract_text)
            if text:
                abstract_parts.append(f"{section_label}: {text}" if section_label else text)
        abstract = "\n".join(abstract_parts)

        publication_types = [
            element_text(publication_type)
            for publication_type in article.findall("PublicationTypeList/PublicationType")
            if element_text(publication_type)
        ]

        doi = None
        pmcid = None
        for article_id in pubmed_article.findall(".//PubmedData/ArticleIdList/ArticleId"):
            id_type = article_id.attrib.get("IdType", "").lower()
            value = element_text(article_id)
            if id_type == "doi":
                doi = value.lower()
            elif id_type == "pmc":
                pmcid = value

        if pmcid:
            access_status = "Full text available through PubMed Central"
        elif abstract:
            access_status = "Abstract available; publisher access may be required"
        else:
            access_status = "Citation only; publisher access may be required"

        records.append(
            {
                "pmid": pmid,
                "title": title,
                "journal": journal,
                "iso_abbreviation": iso_abbreviation,
                "year": year,
                "authors": authors,
                "abstract": abstract,
                "publication_types": publication_types,
                "doi": doi,
                "pmcid": pmcid,
                "pubmed_url": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/" if pmid else "",
                "evidence_type": classify_evidence_type(publication_types, journal),
                "source_group": classify_source_group(journal, iso_abbreviation),
                "access_status": access_status,
            }
        )

    return records


def compare_with_uploaded_inventory(
    records: list[dict[str, Any]],
    inventory: dict[str, Any],
) -> list[dict[str, Any]]:
    existing_records = inventory.get("records", [])
    existing_pmids = {
        str(record.get("pmid")).strip()
        for record in existing_records
        if record.get("pmid")
    }
    existing_dois = {
        str(record.get("doi")).strip().lower()
        for record in existing_records
        if record.get("doi")
    }
    existing_titles = {
        normalize_title(record.get("title"))
        for record in existing_records
        if normalize_title(record.get("title"))
    }
    newest_year = inventory.get("newest_year")

    compared: list[dict[str, Any]] = []
    for record in records:
        record_pmid = str(record.get("pmid") or "").strip()
        record_doi = str(record.get("doi") or "").strip().lower()
        record_title = normalize_title(record.get("title"))

        matched = (
            (record_pmid and record_pmid in existing_pmids)
            or (record_doi and record_doi in existing_dois)
            or (record_title and record_title in existing_titles)
        )

        if matched:
            comparison_status = "Already represented in uploaded literature"
        elif newest_year and record.get("year") and record["year"] > newest_year:
            comparison_status = "Newer potential evidence"
        else:
            comparison_status = "Potentially relevant evidence not matched in uploaded literature"

        compared_record = dict(record)
        compared_record["comparison_status"] = comparison_status
        compared.append(compared_record)

    priority = {
        "Newer potential evidence": 0,
        "Potentially relevant evidence not matched in uploaded literature": 1,
        "Already represented in uploaded literature": 2,
    }

    return sorted(
        compared,
        key=lambda record: (
            priority.get(record.get("comparison_status", ""), 3),
            -(record.get("year") or 0),
            record.get("title") or "",
        ),
    )


def run_evidence_validation_search(
    client: OpenAI,
    email: str,
    guideline_name: str,
    focused_topic: str,
    evidence_question: str,
    literature_text: str,
) -> dict[str, Any]:
    inventory = extract_literature_inventory(client, literature_text)
    base_query = build_pubmed_query(client, guideline_name, focused_topic, evidence_question)
    newest_year = inventory.get("newest_year")

    queries: list[tuple[str, str]] = []
    if newest_year:
        recent_query = (
            f"({base_query}) AND "
            f'(\"{newest_year + 1}/01/01\"[Date - Publication] : '
            '"3000"[Date - Publication])'
        )
        queries.append(("Newer evidence search", recent_query))
    else:
        queries.append(("Recent evidence search", base_query))

    high_value_filter = (
        "systematic review[Publication Type] OR "
        "meta-analysis[Publication Type] OR "
        "randomized controlled trial[Publication Type] OR "
        "guideline[Publication Type] OR "
        "practice guideline[Publication Type] OR "
        '"Cochrane Database Syst Rev"[Journal] OR '
        '"JAMA Pediatr"[Journal] OR '
        '"Pediatrics"[Journal] OR '
        '"Hosp Pediatr"[Journal]'
    )
    queries.append(("High value and targeted source search", f"({base_query}) AND ({high_value_filter})"))

    all_pmids: list[str] = []
    query_details: list[dict[str, str]] = []

    for label, query in queries:
        pmids = pubmed_search(query, email=email, retmax=18)
        query_details.append({"label": label, "query": query, "result_count": str(len(pmids))})
        for pmid in pmids:
            if pmid not in all_pmids:
                all_pmids.append(pmid)

    fetched_records = pubmed_fetch(all_pmids[:30], email=email)
    compared_records = compare_with_uploaded_inventory(fetched_records, inventory)

    return {
        "search_date": datetime.now().strftime("%B %d, %Y"),
        "base_query": base_query,
        "queries": query_details,
        "inventory": inventory,
        "records": compared_records,
    }


def citation_for_prompt(record: dict[str, Any]) -> str:
    authors = record.get("authors") or []
    author_text = ", ".join(authors[:6])
    if len(authors) > 6:
        author_text += ", et al."

    citation_parts = [
        author_text,
        record.get("title") or "",
        record.get("journal") or "",
        str(record.get("year") or ""),
    ]
    citation = ". ".join(part for part in citation_parts if part).strip()

    if record.get("doi"):
        citation += f". DOI: {record['doi']}"
    if record.get("pmid"):
        citation += f". PMID: {record['pmid']}"
    if record.get("pubmed_url"):
        citation += f". {record['pubmed_url']}"

    return citation


def selected_validation_records() -> list[dict[str, Any]]:
    validation = st.session_state.get("validation_results") or {}
    records = validation.get("records", [])
    selected: list[dict[str, Any]] = []

    for record in records:
        pmid = record.get("pmid") or normalize_title(record.get("title"))
        if st.session_state.get(f"select_validation_{pmid}", False):
            selected.append(record)

    return selected


# ============================================================
# Physician Change Card prompt
# ============================================================
PHYSICIAN_CARD_PROMPT = """
You are assisting an Evidence Based Outcomes Center with clinical guideline development and revision.
Generate a physician facing Change Card for clinician and workgroup review.

Core evidence rules
1. Use only the documents, literature evidence, and selected PubMed validation records supplied by the user.
2. Do not invent studies, citations, statistics, quotations, DOIs, PMIDs, links, or source details.
3. Do not use the labels AI, AI suggested, AI recommendation, or evidence packet.
4. Refer to uploaded evidence as literature evidence or supporting literature evidence.
5. Every proposed recommendation, replacement, or new section must be explicitly based on the supplied evidence.
6. Distinguish full text evidence from abstract only evidence. Do not imply that full text was reviewed when only an abstract was supplied.
7. If selected PubMed validation records are provided, use only the records the user selected.

Guideline project rules
8. The project type will be supplied as either New Guideline Development or Revision or Update of Existing Guideline.
9. For New Guideline Development, treat the uploaded guideline as a skeleton or partial template. Draft complete evidence based content under the appropriate headings. Create additional clinically necessary sections when the template does not contain them.
10. For Revision or Update of Existing Guideline, compare the evidence with current guidance. Replace, delete, clarify, expand, or add new sections as required. An update is not limited to sentence replacement.
11. Do not shorten or omit clinically relevant proposed content merely to keep the Change Card brief. Use as many paragraphs, subsections, criteria, or steps as needed to create usable guideline content.
12. When existing wording should be removed or replaced, quote the exact text or identify the paragraph by its beginning and ending words. Do not rely only on line numbers because formatting may change.
13. Always provide concrete proposed wording. Do not ask clinicians to create guidance from scratch.
14. Recommend an addendum only when the content is long, operational, criteria based, stepwise, or likely to clutter the main guideline body.
15. If an addendum is recommended, provide a title, brief main body cross reference language, and complete addendum language.
16. If an addendum is not recommended, use the simplest appropriate placement in the main guideline.

Clinician review rules
17. Frame review as Approve, Modify, or Reject.
18. Do not create unnecessary open ended clinician questions.
19. When a true local decision is required, provide an evidence based recommended default and ask the workgroup to approve or modify it.
20. When no additional local decision is required, state that review is limited to Approve, Modify, or Reject of the proposed wording.
21. Final clinical approval and publishing remain human led.
22. Do not include Proposed Algorithm Impact.
23. Do not include Draft Change Log Entry.
24. If a separate algorithm review may be needed, mention it briefly only in Evidence Traceback Notes.

Confidence rules
25. Label the confidence section exactly GRADE-informed Confidence Rating.
26. The rating must be High, Moderate, or Low.
27. Base the rating on evidence type, consistency, directness, precision, pediatric applicability, and local implementation considerations.
28. Begin the section with the exact format **Rating:** High, **Rating:** Moderate, or **Rating:** Low.
29. Provide a concise rationale for the rating.

Citation organization rules
30. Evidence Traceback Notes must include every relevant source used in the proposed change.
31. Preserve complete citations, DOI, PMID, and hyperlinks when those details are present in the supplied evidence.
32. Organize citations under these headings when applicable, and omit empty headings:
### Clinical Guidelines and Consensus Statements
### Systematic Reviews and Meta Analyses
### Randomized Trials
### Observational Studies
### Other Relevant Evidence

Section 5 formatting rules
33. Include these bold labels in Section 5:
**Recommended update action:**
**Suggested guideline section:**
**Exact placement:**
**Existing wording to delete or replace:**
**Proposed guideline wording:**
**Related alignment notes:**
34. For New Guideline Development, write Not applicable under Existing wording to delete or replace when there is no existing content.
35. Section 5 may contain multiple proposed sections when the evidence question requires them.

Output exactly these main sections:
# Change Card Title
## 1. Evidence Question
## 2. Current Guidance
## 3. Workgroup Concern / Reason for Review
## 4. Evidence Finding
## 5. Proposed Guideline Edit or New Section
## 6. GRADE-informed Confidence Rating
## 7. Decision Required Before Approval
## 8. Evidence Traceback Notes
"""


# ============================================================
# Confidence display
# ============================================================
def extract_confidence_level(card_text: str) -> str:
    section_match = re.search(
        r"##\s*6\.\s*GRADE[- ]informed Confidence Rating(.*?)(?=\n##\s*7\.|\Z)",
        card_text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    text_to_search = section_match.group(1) if section_match else card_text

    for level in ("High", "Moderate", "Low"):
        if re.search(rf"\b{level}\b", text_to_search, flags=re.IGNORECASE):
            return level
    return "Not identified"


def render_confidence_badge(level: str) -> None:
    label, _, hex_color = confidence_color(level)
    background_map = {
        "High": "#edf7ed",
        "Moderate": "#fff8e1",
        "Low": "#fdecec",
        "Not identified": "#f1f5f9",
    }
    text_map = {
        "High": "#1b5e20",
        "Moderate": "#7c4a03",
        "Low": "#8e1b1b",
        "Not identified": "#475569",
    }

    st.markdown(
        f"""
        <div class="confidence-badge" style="background:{background_map[level]};
        color:{text_map[level]};border-color:#{hex_color};">
            <span class="confidence-dot" style="background:#{hex_color};"></span>
            GRADE informed confidence: {level}
            <span style="font-weight:500;">Stoplight: {label}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# Compact status chips
# ============================================================
def render_status_chip(label: str, status: str) -> None:
    class_name = {
        "ready": "status-ready",
        "missing": "status-missing",
        "neutral": "status-neutral",
    }[status]

    symbol = {
        "ready": "✓",
        "missing": "!",
        "neutral": "○",
    }[status]

    st.markdown(
        f'<span class="status-chip {class_name}">{symbol} {label}</span>',
        unsafe_allow_html=True,
    )


# ============================================================
# App layout
# ============================================================
render_branding()

st.write(
    "Generate an evidence based physician facing Change Card for guideline development or revision review."
)
st.info(
    "This prototype is intended for non PHI guideline materials only. Final clinical approval remains human led."
)

api_key = get_api_key()
ncbi_email = get_secret("NCBI_EMAIL")

with st.sidebar:
    st.header("Status")
    if api_key:
        st.success("OpenAI key available")
    else:
        st.error("OpenAI key missing")

    if ncbi_email:
        st.success("NCBI contact email available")
    else:
        st.info("NCBI contact email not configured")

    st.markdown("### Section navigation")
    st.markdown(
        """
        [1. Guideline details](#1-guideline-update-details)  
        [2. Upload documents](#2-upload-documents)  
        [3. Evidence validation](#3-evidence-validation)  
        [4. Readiness check](#4-readiness-check)  
        [5. Generate Change Card](#5-generate-change-card)  
        [6. Preview and download](#6-preview-and-download)
        """
    )

    st.caption(
        "The hosted app should use Streamlit Secrets. Do not place credentials in GitHub."
    )


# ============================================================
# Section 1: Guideline details
# ============================================================
st.header("1. Guideline Update Details")

project_type = st.radio(
    "Guideline project type",
    options=PROJECT_TYPES,
    horizontal=True,
    help=(
        "Choose whether the app should revise existing guidance or develop content "
        "within a new guideline skeleton."
    ),
)

left_col, right_col = st.columns(2)
with left_col:
    guideline_name = st.text_input(
        "Guideline name (required)",
        value="",
        placeholder="Example: Bronchiolitis",
    )
    focused_topic = st.text_input(
        "Focused topic (required)",
        value="",
        placeholder="Example: Feeding, NG Feeds, and IV Fluids",
    )

with right_col:
    if project_type == "New Guideline Development":
        date_label = "Guideline template or project date (required)"
        date_help = "Enter the template date or the date this guideline development project began."
    else:
        date_label = "Guideline or pathway last updated date (required)"
        date_help = "Enter the date shown on the current guideline or pathway."

    guideline_date = st.text_input(
        date_label,
        value="",
        placeholder="Example: October 8, 2019",
        help=date_help,
    )

st.write("Evidence question")
evidence_question = st.text_area(
    "Paste the question being answered (required)",
    value="",
    height=145,
    placeholder="Paste the PICO question or focused evidence question here.",
)


# ============================================================
# Section 2: Upload documents
# ============================================================
st.header("2. Upload Documents")
st.write(
    "Upload the guideline or skeleton and the primary literature evidence. Supporting local documents are optional."
)

upload_left, upload_right = st.columns(2)
with upload_left:
    guideline_label = (
        "Guideline skeleton template (required, DOCX or PDF)"
        if project_type == "New Guideline Development"
        else "Current guideline draft (required, DOCX or PDF)"
    )
    guideline_file = st.file_uploader(
        guideline_label,
        type=["docx", "pdf"],
        key="guideline_file",
    )

    pathway_file = st.file_uploader(
        "Current algorithm or pathway (optional, DOCX or PDF)",
        type=["docx", "pdf"],
        key="pathway_file",
    )

    notes_file = st.file_uploader(
        "Workgroup notes or meeting minutes (optional, DOCX or PDF)",
        type=["docx", "pdf"],
        key="notes_file",
    )

with upload_right:
    literature_file = st.file_uploader(
        "Literature evidence from OpenEvidence or another source (required, DOCX or PDF)",
        type=["docx", "pdf"],
        key="literature_file",
    )

    optional_context_file = st.file_uploader(
        "Local pathway or context document (optional, DOCX or PDF)",
        type=["docx", "pdf"],
        key="optional_context_file",
    )


# ============================================================
# Section 3: Evidence validation
# ============================================================
current_validation_signature = validation_context_signature(
    guideline_name=guideline_name,
    focused_topic=focused_topic,
    evidence_question=evidence_question,
    literature_file=literature_file,
)

stored_validation = st.session_state.get("validation_results")
if (
    stored_validation
    and stored_validation.get("context_signature") != current_validation_signature
):
    st.session_state.pop("validation_results", None)
    for session_key in list(st.session_state.keys()):
        if str(session_key).startswith("select_validation_"):
            st.session_state.pop(session_key, None)

st.header("3. Evidence Validation")
run_validation = st.toggle(
    "Run Evidence Validation Search",
    value=False,
    help=(
        "Search PubMed and MEDLINE for potentially newer or missing evidence, including "
        "indexed JAMA Pediatrics, AAP, and Cochrane records."
    ),
)

if run_validation:
    st.caption(
        "This validation search compares PubMed indexed records with citations recognized in the uploaded literature. "
        "It does not guarantee that every relevant study has been found."
    )

    validation_prerequisites = all(
        [
            api_key,
            ncbi_email,
            guideline_name.strip(),
            focused_topic.strip(),
            evidence_question.strip(),
            literature_file is not None,
        ]
    )

    if not ncbi_email:
        st.warning("Add NCBI_EMAIL to Streamlit Secrets before running the validation search.")

    validation_button = st.button(
        "Run PubMed Validation Search",
        disabled=not validation_prerequisites,
        type="secondary",
    )

    if validation_button:
        try:
            client = create_openai_client(api_key)
            with st.status("Running evidence validation...", expanded=True) as status:
                st.write("Reading the uploaded literature evidence")
                literature_text_for_validation = read_uploaded_file(literature_file)
                if not literature_text_for_validation.strip():
                    raise ValueError("No readable text was extracted from the literature evidence file.")

                st.write("Extracting citation fingerprints from the uploaded literature")
                st.write("Building a focused PubMed query")
                st.write("Searching PubMed and MEDLINE")
                st.write("Comparing search results with the uploaded citations")

                validation_results = run_evidence_validation_search(
                    client=client,
                    email=ncbi_email,
                    guideline_name=guideline_name,
                    focused_topic=focused_topic,
                    evidence_question=evidence_question,
                    literature_text=literature_text_for_validation,
                )

                validation_results["context_signature"] = current_validation_signature
                st.session_state["validation_results"] = validation_results
                status.update(label="Evidence validation completed", state="complete")
        except Exception as error:
            st.error("The evidence validation search could not be completed.")
            st.code(str(error))

    validation_results = st.session_state.get("validation_results")
    if validation_results:
        records = validation_results.get("records", [])
        inventory = validation_results.get("inventory", {})
        newest_year = inventory.get("newest_year")

        newer_count = sum(
            record.get("comparison_status") == "Newer potential evidence"
            for record in records
        )
        unmatched_count = sum(
            record.get("comparison_status")
            == "Potentially relevant evidence not matched in uploaded literature"
            for record in records
        )
        represented_count = sum(
            record.get("comparison_status")
            == "Already represented in uploaded literature"
            for record in records
        )

        st.info(
            "Validation summary: "
            f"newest citation year recognized in the uploaded literature: "
            f"{newest_year or 'not identified'}; "
            f"newer candidates: {newer_count}; "
            f"other unmatched candidates: {unmatched_count}; "
            f"records already represented: {represented_count}."
        )

        with st.expander("View documented PubMed search strategy", expanded=False):
            st.write(f"Search date: {validation_results.get('search_date', '')}")
            for query_info in validation_results.get("queries", []):
                st.markdown(f"**{query_info.get('label', 'Search')}**")
                st.code(query_info.get("query", ""), language=None)

        additional_records = [
            record
            for record in records
            if record.get("comparison_status") != "Already represented in uploaded literature"
        ]

        if not additional_records:
            st.success("No additional candidate records were identified by this documented search.")
        else:
            st.subheader("Candidate evidence for your review")
            st.caption(
                "Select only the records you want the Change Card to consider. Abstract only records remain labeled as such."
            )

            for record in additional_records:
                record_key = record.get("pmid") or normalize_title(record.get("title"))
                title = record.get("title") or "Untitled PubMed record"
                year = record.get("year") or "Date unavailable"
                journal = record.get("journal") or "Journal unavailable"
                source_group = record.get("source_group") or "PubMed indexed source"

                st.markdown('<div class="validation-card">', unsafe_allow_html=True)
                st.checkbox(
                    f"Include: {title}",
                    key=f"select_validation_{record_key}",
                )
                st.markdown(
                    f"**{record.get('comparison_status')}**  \n"
                    f"{journal}, {year}  \n"
                    f"Source group: {source_group}  \n"
                    f"Evidence type: {record.get('evidence_type')}  \n"
                    f"Access: {record.get('access_status')}"
                )
                if record.get("pubmed_url"):
                    st.markdown(f"[Open in PubMed]({record['pubmed_url']})")
                with st.expander("View abstract and citation", expanded=False):
                    st.write(citation_for_prompt(record))
                    st.write(record.get("abstract") or "No abstract was available through PubMed.")
                st.markdown("</div>", unsafe_allow_html=True)
else:
    st.info(
        "Evidence validation is optional. When it is off, the app uses the uploaded literature evidence only."
    )


# ============================================================
# Section 4: Readiness check
# ============================================================
st.header("4. Readiness Check")

stored_validation_results = st.session_state.get("validation_results")
validation_complete = (
    not run_validation
    or (
        stored_validation_results is not None
        and stored_validation_results.get("context_signature")
        == current_validation_signature
    )
)

required_items = [
    ("OpenAI key", bool(api_key)),
    ("Project type", bool(project_type)),
    ("Guideline name", bool(guideline_name.strip())),
    ("Focused topic", bool(focused_topic.strip())),
    ("Evidence question", bool(evidence_question.strip())),
    ("Guideline date", bool(guideline_date.strip())),
    ("Guideline or skeleton", guideline_file is not None),
    ("Literature evidence", literature_file is not None),
]

if run_validation:
    required_items.append(("NCBI contact email", bool(ncbi_email)))
    required_items.append(("Validation search", validation_complete))

optional_items = [
    ("Algorithm or pathway", pathway_file is not None),
    ("Workgroup notes", notes_file is not None),
    ("Local context", optional_context_file is not None),
]

all_items = [(label, ready, False) for label, ready in required_items]
all_items.extend((label, uploaded, True) for label, uploaded in optional_items)

for row_start in range(0, len(all_items), 3):
    row_items = all_items[row_start:row_start + 3]
    columns = st.columns(3)
    for column, (label, ready_value, optional) in zip(columns, row_items):
        with column:
            if optional:
                render_status_chip(
                    f"{label}: {'uploaded' if ready_value else 'not uploaded'}",
                    "ready" if ready_value else "neutral",
                )
            else:
                render_status_chip(
                    f"{label}: {'ready' if ready_value else 'missing'}",
                    "ready" if ready_value else "missing",
                )

ready = all(ready_value for _, ready_value in required_items)


# ============================================================
# Section 5: Generate Change Card
# ============================================================
st.header("5. Generate Change Card")

if not ready:
    st.warning("Complete the required fields and uploads before generating the Change Card.")

selected_records = selected_validation_records() if run_validation else []

if run_validation and st.session_state.get("validation_results"):
    st.caption(f"Selected validation records: {len(selected_records)}")

generate = st.button(
    "Generate Physician Facing Change Card",
    disabled=not ready,
    type="primary",
)

if generate:
    try:
        client = create_openai_client(api_key)

        progress = st.progress(0)
        with st.status("Preparing the Change Card...", expanded=True) as status:
            st.write("Validating required information")
            progress.progress(10)

            st.write("Reading the guideline or skeleton")
            guideline_text = read_uploaded_file(guideline_file)
            if not guideline_text.strip():
                raise ValueError("No readable text was extracted from the guideline or skeleton.")
            progress.progress(25)

            st.write("Reading literature evidence")
            literature_evidence_text = read_uploaded_file(literature_file)
            if not literature_evidence_text.strip():
                raise ValueError("No readable text was extracted from the literature evidence.")
            progress.progress(40)

            st.write("Reading optional supporting documents")
            pathway_text = read_uploaded_file(pathway_file) if pathway_file else ""
            notes_text = read_uploaded_file(notes_file) if notes_file else ""
            optional_context_text = (
                read_uploaded_file(optional_context_file)
                if optional_context_file
                else ""
            )
            progress.progress(55)

            validation_text = ""
            if selected_records:
                selected_sections: list[str] = []
                for index, record in enumerate(selected_records, start=1):
                    selected_sections.append(
                        f"""
SELECTED VALIDATION RECORD {index}
Citation: {citation_for_prompt(record)}
Evidence type: {record.get('evidence_type')}
Access reviewed: {record.get('access_status')}
Abstract supplied by PubMed:
{record.get('abstract') or 'No abstract available.'}
"""
                    )
                validation_text = "\n".join(selected_sections)

            st.write("Building the evidence grounded generation request")
            user_input = f"""
GUIDELINE PROJECT TYPE:
{project_type}

GUIDELINE NAME:
{guideline_name}

FOCUSED TOPIC:
{focused_topic}

EVIDENCE QUESTION:
{evidence_question}

GUIDELINE OR TEMPLATE DATE:
{guideline_date}

CURRENT GUIDELINE OR SKELETON TEXT:
{guideline_text}

LITERATURE EVIDENCE:
{literature_evidence_text}

CURRENT ALGORITHM OR PATHWAY TEXT, IF PROVIDED:
{pathway_text or 'Not provided.'}

WORKGROUP NOTES OR MEETING MINUTES, IF PROVIDED:
{notes_text or 'Not provided.'}

LOCAL PATHWAY OR CONTEXT DOCUMENT, IF PROVIDED:
{optional_context_text or 'Not provided.'}

USER SELECTED PUBMED VALIDATION RECORDS, IF ANY:
{validation_text or 'No additional PubMed validation records were selected.'}
"""
            progress.progress(70)

            st.write("Drafting the physician facing Change Card")
            response = client.responses.create(
                model=MODEL_NAME,
                instructions=PHYSICIAN_CARD_PROMPT,
                input=user_input,
                store=False,
            )
            physician_card = response.output_text
            progress.progress(90)

            st.write("Preparing the preview and Word document")
            confidence_level = extract_confidence_level(physician_card)
            st.session_state["physician_card"] = physician_card
            st.session_state["focused_topic"] = focused_topic
            st.session_state["guideline_name"] = guideline_name
            st.session_state["confidence_level"] = confidence_level
            progress.progress(100)

            status.update(label="Change Card generated successfully", state="complete")

    except Exception as error:
        st.error("The Change Card could not be generated.")
        st.write("Error details")
        st.code(str(error))


# ============================================================
# Section 6: Preview and download
# ============================================================
if "physician_card" in st.session_state:
    st.header("6. Preview and Download")

    physician_card = st.session_state["physician_card"]
    topic_for_file = st.session_state.get("focused_topic", "Change Card")
    guideline_for_file = st.session_state.get("guideline_name", "Guideline")
    confidence_level = st.session_state.get("confidence_level", "Not identified")

    st.subheader("Preview")
    render_confidence_badge(confidence_level)
    st.markdown(physician_card)

    docx_file = create_styled_docx(
        physician_card,
        guideline_name=guideline_for_file,
        topic=topic_for_file,
        confidence_level=confidence_level,
    )

    safe_guideline = re.sub(r"[^A-Za-z0-9]+", "_", guideline_for_file).strip("_")
    safe_topic = re.sub(r"[^A-Za-z0-9]+", "_", topic_for_file).strip("_")
    safe_guideline = safe_guideline or "Guideline"
    safe_topic = safe_topic or "Change_Card"

    filename = f"{safe_guideline}_{safe_topic}_Physician_Facing_Change_Card.docx"

    st.download_button(
        label="Download Word Document",
        data=docx_file,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    st.caption(
        "Review the generated card before sharing. Final guideline approval and publishing remain human led."
    )
